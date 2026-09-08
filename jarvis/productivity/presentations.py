import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


class PresentationBuilder:
    BACKGROUND = RGBColor(5, 12, 20)
    PANEL = RGBColor(10, 25, 38)
    ACCENT = RGBColor(23, 214, 255)
    GOLD = RGBColor(255, 171, 64)
    TEXT = RGBColor(232, 246, 255)
    MUTED = RGBColor(142, 177, 194)

    def __init__(self, output_dir: Path):
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _parse_outline(self, outline: str):
        sections = re.split(r"\n\s*---+\s*\n", outline.strip())
        slides = []
        for section in sections:
            lines = [line.strip() for line in section.splitlines() if line.strip()]
            if not lines:
                continue
            title = lines[0].lstrip("# ").strip()
            bullets = [
                line.lstrip("-*•0123456789. ").strip()
                for line in lines[1:]
                if line.lstrip("-*•0123456789. ").strip()
            ]
            slides.append((title[:100], bullets[:10]))
        if not slides:
            raise ValueError("The presentation outline must contain at least one slide")
        return slides[:20]

    def _set_background(self, slide):
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = self.BACKGROUND
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(0.12),
            Inches(7.5),
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = self.ACCENT
        bar.line.fill.background()

    def _add_footer(self, slide, index: int):
        footer = slide.shapes.add_textbox(
            Inches(0.65),
            Inches(7.08),
            Inches(12.0),
            Inches(0.22),
        )
        paragraph = footer.text_frame.paragraphs[0]
        paragraph.text = f"JARVIS GENERATED  •  {index:02d}"
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(9)
        paragraph.font.color.rgb = self.MUTED
        paragraph.alignment = PP_ALIGN.RIGHT

    def _add_title_slide(self, deck, title: str, subtitle: str):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        self._set_background(slide)
        accent = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(9.75),
            Inches(1.0),
            Inches(2.25),
            Inches(2.25),
        )
        accent.fill.background()
        accent.line.color.rgb = self.GOLD
        accent.line.width = Pt(3)

        title_box = slide.shapes.add_textbox(
            Inches(0.9),
            Inches(2.25),
            Inches(10.8),
            Inches(1.6),
        )
        title_paragraph = title_box.text_frame.paragraphs[0]
        title_paragraph.text = title
        title_paragraph.font.name = "Aptos Display"
        title_paragraph.font.size = Pt(34)
        title_paragraph.font.bold = True
        title_paragraph.font.color.rgb = self.TEXT

        subtitle_box = slide.shapes.add_textbox(
            Inches(0.95),
            Inches(4.05),
            Inches(9.5),
            Inches(0.7),
        )
        subtitle_paragraph = subtitle_box.text_frame.paragraphs[0]
        subtitle_paragraph.text = subtitle or "Prepared with JARVIS"
        subtitle_paragraph.font.name = "Aptos"
        subtitle_paragraph.font.size = Pt(18)
        subtitle_paragraph.font.color.rgb = self.ACCENT
        self._add_footer(slide, 1)

    def _add_content_slide(
        self,
        deck,
        title: str,
        bullets: list[str],
        index: int,
    ):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        self._set_background(slide)

        title_box = slide.shapes.add_textbox(
            Inches(0.75),
            Inches(0.55),
            Inches(11.8),
            Inches(0.75),
        )
        title_paragraph = title_box.text_frame.paragraphs[0]
        title_paragraph.text = title
        title_paragraph.font.name = "Aptos Display"
        title_paragraph.font.size = Pt(27)
        title_paragraph.font.bold = True
        title_paragraph.font.color.rgb = self.TEXT

        panel = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.75),
            Inches(1.55),
            Inches(11.8),
            Inches(5.15),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = self.PANEL
        panel.line.color.rgb = self.ACCENT
        panel.line.transparency = 60

        text_box = slide.shapes.add_textbox(
            Inches(1.15),
            Inches(1.9),
            Inches(11.0),
            Inches(4.45),
        )
        frame = text_box.text_frame
        frame.word_wrap = True
        frame.clear()
        content = bullets or ["Add supporting information for this section."]
        for bullet_index, bullet in enumerate(content):
            paragraph = (
                frame.paragraphs[0]
                if bullet_index == 0
                else frame.add_paragraph()
            )
            paragraph.text = bullet
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(20)
            paragraph.font.color.rgb = self.TEXT
            paragraph.space_after = Pt(13)
            paragraph.level = 0
            paragraph.text = f"•  {paragraph.text}"
        self._add_footer(slide, index)

    def create(
        self,
        title: str,
        outline: str,
        subtitle: str = "",
        filename: str = "",
    ) -> str:
        normalized_title = title.strip()
        if not normalized_title:
            raise ValueError("Presentation title cannot be empty")
        slides = self._parse_outline(outline)
        deck = Presentation()
        deck.slide_width = Inches(13.333)
        deck.slide_height = Inches(7.5)
        self._add_title_slide(deck, normalized_title, subtitle.strip())
        for index, (slide_title, bullets) in enumerate(slides, start=2):
            self._add_content_slide(deck, slide_title, bullets, index)

        requested_name = Path(filename).stem if filename.strip() else normalized_title
        safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", requested_name).strip("_")
        output = self.output_dir / f"{safe_name[:80] or 'presentation'}.pptx"
        deck.save(output)
        return str(output.resolve())
